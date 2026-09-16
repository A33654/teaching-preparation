"""教案导出为 DOCX"""

import io
from datetime import datetime
from docx import Document as DocxDocument
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH


def _format_date(val) -> str:
    """安全格式化日期"""
    if val is None:
        return ""
    if isinstance(val, datetime):
        return val.strftime("%Y-%m-%d")
    if isinstance(val, str):
        return val[:10] if len(val) >= 10 else val
    return str(val)[:10]


def export_lesson_plan_to_docx(lesson_plan: dict) -> bytes:
    """将教案导出为 .docx 文件，返回二进制内容"""
    doc = DocxDocument()

    # 页面设置
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.5)
    section.bottom_margin = Cm(2.5)

    # 标题
    title = doc.add_heading(lesson_plan.get("title", "教案"), level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # 基本信息
    info_table = doc.add_table(rows=3, cols=4)
    info_table.style = "Light Grid Accent 1"
    info_cells = [
        ("学科", lesson_plan.get("subject", "")),
        ("学段", lesson_plan.get("grade_level", "")),
        ("创建时间", _format_date(lesson_plan.get("created_at"))),
        ("", ""),
    ]
    for i, (label, value) in enumerate(info_cells[:2]):
        row = info_table.rows[i]
        row.cells[0].text = label
        row.cells[1].text = str(value)
        row.cells[2].text = info_cells[i + 2][0] if i + 2 < len(info_cells) else ""
        row.cells[3].text = str(info_cells[i + 2][1]) if i + 2 < len(info_cells) else ""

    doc.add_paragraph()

    # 关联知识点
    kps = lesson_plan.get("knowledge_points", [])
    if kps:
        doc.add_heading("关联知识点", level=2)
        kp_text = "、".join([kp.get("name", "") for kp in kps if kp.get("name")])
        doc.add_paragraph(kp_text or "无")

    # 教学目标
    objectives = lesson_plan.get("teaching_objectives", "")
    if objectives:
        doc.add_heading("教学目标", level=2)
        for line in objectives.replace("；", "\n").replace(";", "\n").split("\n"):
            line = line.strip()
            if line:
                doc.add_paragraph(line, style="List Bullet")

    # 教学重点
    key_points = lesson_plan.get("key_points", [])
    if key_points:
        doc.add_heading("教学重点", level=2)
        for pt in key_points:
            doc.add_paragraph(str(pt), style="List Bullet")

    # 教学难点
    difficult_points = lesson_plan.get("difficult_points", [])
    if difficult_points:
        doc.add_heading("教学难点", level=2)
        for pt in difficult_points:
            doc.add_paragraph(str(pt), style="List Bullet")

    # 教学过程
    process = lesson_plan.get("teaching_process", "")
    if process:
        doc.add_heading("教学过程", level=2)
        for line in process.replace("\r\n", "\n").split("\n"):
            line = line.strip()
            if not line:
                continue
            if line.startswith(("一", "二", "三", "四", "五", "六", "七", "八", "九", "十")):
                doc.add_heading(line, level=3)
            else:
                doc.add_paragraph(line)

    # 保存到内存
    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer.read()
